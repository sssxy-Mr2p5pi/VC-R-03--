"""Read-only exploration of whole-invoice sample starts, not app acceptance."""
from pathlib import Path
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from bisect import bisect_right
from time import perf_counter
import hashlib
import json
import sys

ROOT = Path(r'E:\佘逍遥的大学_大二\PythonCoding\VC-R-03-退货对账与差异处理系统')
sys.path.insert(0, str(ROOT/'scripts'))
from batch_validate_rules import norm, nonproduct
from openpyxl import load_workbook

source = ROOT/'data/raw/online_retail_II.xlsx'
output = Path(__file__).parent/'performance_start_exploration.json'
hash_before = hashlib.sha256(source.read_bytes()).hexdigest()
assert hash_before == 'bcbe73b35f5b7babf197fb0cb983a11f5d9ff929078d4aa53d171b1f2df2e980'
start = perf_counter()
wb = load_workbook(source, read_only=True, data_only=False)
sheets=[]
for ws in wb.worksheets:
    groups=defaultdict(list)
    rows=ws.iter_rows(values_only=True)
    header=next(rows)
    assert tuple(header)==('Invoice','StockCode','Description','Quantity','InvoiceDate','Price','Customer ID','Country')
    count=0
    for rn,row in enumerate(rows,2):
        groups[norm(row[0])].append((rn,row))
        count+=1
    sheets.append((ws.title,count,groups))
    print(json.dumps({'read_sheet':ws.title,'rows':count},ensure_ascii=False),flush=True)
wb.close()

def select(percent):
    selected=[]
    details=[]
    for si,(name,count,groups) in enumerate(sheets):
        threshold=2+int(count*percent)
        remaining=5000
        chosen=[]
        for inv,lines in groups.items():
            if lines[0][0]>=threshold and len(lines)<=remaining:
                chosen.extend((si,rn,row) for rn,row in lines)
                remaining-=len(lines)
                if not remaining:
                    break
        if remaining:
            return None,{'percent':percent,'sheet':name,'remaining':remaining}
        chosen.sort(key=lambda item:item[1])
        selected.extend(chosen)
        details.append({'sheet':name,'threshold_excel_row':threshold,'selected_rows':len(chosen),
                        'first_excel_row':chosen[0][1],'last_excel_row':chosen[-1][1],
                        'invoice_count':len({norm(x[2][0]) for x in chosen})})
    return selected,details

def analyze(selected):
    invoice_sides=[defaultdict(Counter),defaultdict(Counter)]
    for si,rn,row in selected:
        invoice_sides[si][norm(row[0])][row[1:]]+=1
    equal={key for key in invoice_sides[0].keys()&invoice_sides[1].keys()
           if invoice_sides[0][key]==invoice_sides[1][key]}
    kept=[item for item in selected if not(item[0]==0 and norm(item[2][0]) in equal)
          and not(not norm(item[2][0]).startswith('C') and item[2][3]<0)]
    originals=defaultdict(list)
    cancels=[]
    np=0
    for si,rn,row in kept:
        inv=norm(row[0])
        if nonproduct(row):
            np+=1
        elif inv.startswith('C'):
            cancels.append((si,rn,row))
        elif row[3]>0 and row[6] is not None and isinstance(row[4],datetime):
            originals[(norm(row[6]),norm(row[1]))].append((row[4],si,rn,row))
    for lines in originals.values():
        lines.sort(key=lambda x:(x[0],x[1],x[2]))
    dates={key:[x[0] for x in lines] for key,lines in originals.items()}
    counts=Counter()
    claims=defaultdict(list)
    valid_price=set()
    for i,(si,rn,row) in enumerate(cancels):
        if row[6] is None:
            counts['missing_customer']+=1
            continue
        key=(norm(row[6]),norm(row[1]))
        candidates=originals.get(key,[])[:bisect_right(dates.get(key,[]),row[4])]
        if not candidates:
            counts['no_candidate']+=1
        elif len(candidates)>1:
            counts['multiple_candidates']+=1
        else:
            counts['unique_candidate']+=1
            o=candidates[0]
            oid=(o[1],o[2])
            claims[oid].append((i,-row[3],o[3][3]))
            if Decimal(str(row[5]))==Decimal(str(o[3][5])) and row[3]<0 and row[5]>0:
                valid_price.add(i)
            else:
                counts['unique_price_or_quantity_difference']+=1
    over={i for claim in claims.values() if sum(q for i,q,orig in claim)>claim[0][2] for i,q,orig in claim}
    return {'raw_rows':len(selected),'retained_rows':len(kept),'excluded_rows':len(selected)-len(kept),
            'nonproduct_retained_rows':np,'merchandise_cancellation_rows':len(cancels),
            'with_candidates':counts['unique_candidate']+counts['multiple_candidates'],
            'proposed_auto_rows':len(valid_price-over),'cumulative_over_quota_rows':len(over),
            'candidate_counts':dict(counts)}

cases=[]
for percent in [0,0.25,0.5,0.75]:
    selected,details=select(percent)
    cases.append({'percent':percent,'selection':details,'activity':analyze(selected) if selected else None})
assert hashlib.sha256(source.read_bytes()).hexdigest()==hash_before
report={'measured_at':datetime.now(timezone(timedelta(hours=8))).isoformat(),
        'source_sha256':hash_before,'source_unchanged':True,
        'scope':'Both sheets, 5000 raw rows each, whole-invoice greedy selection. Four start positions explored openly. Candidate activity uses existing design dry-run rule; not app/performance acceptance.',
        'cases':cases,'elapsed_seconds':perf_counter()-start}
output.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2),flush=True)
