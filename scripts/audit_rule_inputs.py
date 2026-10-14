"""Read-only evidence for the VC-R-03 rule review; no business matching is run."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import time
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

EXACT = {
    "POST", "DOT", "M", "D", "C2", "S", "B", "PADS", "ADJUST", "ADJUST2",
    "BANK CHARGES", "AMAZONFEE", "CRUK",
}


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    if args.source.resolve() == args.output.resolve():
        raise ValueError("Source and output must differ")
    started = time.perf_counter()
    before = sha(args.source)
    wb = load_workbook(args.source, read_only=True, data_only=True)
    first: dict[str, Counter] = defaultdict(Counter)
    first_stats: dict[str, Counter] = defaultdict(Counter)
    overlap_second: dict[str, Counter] = defaultdict(Counter)
    variants: dict[str, set[str]] = defaultdict(set)
    reports = {}
    for sheet_index, ws in enumerate(wb.worksheets):
        rows = ws.iter_rows(values_only=True)
        header = next(rows)
        counts = Counter()
        scales = Counter()
        examples = {}
        invoices: dict[str, set[str]] = defaultdict(set)
        for excel_row, row in enumerate(rows, 2):
            inv, sku, _desc, qty, _date, price, customer, _country = row
            invoice = str(inv).strip().upper()
            raw_code = "" if sku is None else str(sku)
            code = unicodedata.normalize("NFKC", raw_code).strip().upper()
            variants[code].add(raw_code)
            nonproduct = (
                invoice.startswith("A") or code in EXACT
                or bool(re.fullmatch(r"TEST\d+|GIFT_\d+_\d+", code))
            )
            cancel = invoice.startswith("C")
            negative_regular = not cancel and qty < 0
            flags = {
                "rows": True,
                "nonproduct": nonproduct,
                "negative_regular": negative_regular,
                "negative_regular_and_nonproduct": negative_regular and nonproduct,
                "nonproduct_regular": nonproduct and not cancel,
                "nonproduct_cancel": nonproduct and cancel,
                "regular_POST": not cancel and code == "POST",
                "cancel_rows": cancel,
                "commodity_cancel_rows": cancel and not nonproduct,
                "cancel_missing_customer_rows": cancel and customer is None,
                "commodity_cancel_missing_customer_rows": cancel and not nonproduct and customer is None,
                "sku_strip_changes": raw_code != raw_code.strip(),
                "sku_case_changes": raw_code.strip() != raw_code.strip().upper(),
                "sku_nfkc_changes": raw_code != unicodedata.normalize("NFKC", raw_code),
                "sku_missing": not code,
                "sku_leading_zero": bool(re.match(r"^0\d", code)),
                "commodity_price_zero": not nonproduct and price == 0,
                "commodity_price_negative": not nonproduct and price < 0,
                "cancel_commodity_price_nonpositive": cancel and not nonproduct and price <= 0,
                "cancel_commodity_quantity_nonnegative": cancel and not nonproduct and qty >= 0,
            }
            for key, value in flags.items():
                counts[key] += int(value)
                if value and key in {"sku_strip_changes", "sku_case_changes", "sku_nfkc_changes"}:
                    examples.setdefault(key, {"sheet": ws.title, "excel_row": excel_row, "raw": raw_code, "normalized": code})
            places = max(0, -Decimal(str(price)).as_tuple().exponent)
            scales[str(places)] += 1
            if places > 2:
                examples.setdefault("price_more_than_2_decimals", {"sheet": ws.title, "excel_row": excel_row, "price": str(price), "stock_code": raw_code})
            if cancel:
                invoices["cancel"].add(invoice)
                invoices["nonproduct" if nonproduct else "commodity"].add(invoice)
                if customer is not None:
                    invoices["with_customer"].add(invoice)
                    if not nonproduct:
                        invoices["commodity_with_customer"].add(invoice)
            if sheet_index == 0:
                # Exact seven-field multiset, preserving repeated lines within an invoice.
                first[invoice][tuple(row[1:])] += 1
                first_stats[invoice]["rows"] += 1
                first_stats[invoice]["negative_regular"] += int(negative_regular)
                first_stats[invoice]["nonproduct"] += int(nonproduct)
            elif invoice in first:
                overlap_second[invoice][tuple(row[1:])] += 1
        invoice_counts = {key: len(value) for key, value in invoices.items()}
        invoice_counts["pure_nonproduct"] = len(invoices["nonproduct"] - invoices["commodity"])
        invoice_counts["mixed"] = len(invoices["nonproduct"] & invoices["commodity"])
        reports[ws.title] = {
            "header": header, "counts_before_dedup": dict(counts),
            "invoice_counts_before_dedup": invoice_counts,
            "price_decimal_places_before_dedup": dict(scales), "examples": examples,
        }
    wb.close()
    equal = [inv for inv, rows in overlap_second.items() if rows == first[inv]]
    unequal = [inv for inv, rows in overlap_second.items() if rows != first[inv]]
    removed = sum(first_stats[inv]["rows"] for inv in equal)
    overlap_negative = sum(first_stats[inv]["negative_regular"] for inv in equal)
    overlap_nonproduct = sum(first_stats[inv]["nonproduct"] for inv in equal)
    total = sum(v["counts_before_dedup"]["rows"] for v in reports.values())
    negative = sum(v["counts_before_dedup"]["negative_regular"] for v in reports.values())
    excluded = removed + negative - overlap_negative
    after = sha(args.source)
    assert before == after, "Source changed during audit"
    assert total == 1067371, "Unexpected raw row count"
    collisions = {key: sorted(value) for key, value in variants.items() if len(value) > 1}
    result = {
        "measured_at": datetime.now(timezone(timedelta(hours=8))).isoformat(),
        "scope": "Both worksheets; each first row is a header. This is a read-only input audit, not matching or business acceptance.",
        "source": str(args.source.resolve()), "source_sha256": before, "source_unchanged": True,
        "audit_script_sha256": sha(Path(__file__)),
        "nonproduct_rule": {"invoice_prefix": "A", "exact_codes": sorted(EXACT), "fullmatch_patterns": [r"TEST\d+", r"GIFT_\d+_\d+"], "code_normalization_for_audit": "NFKC, strip ends, upper", "status": "Project proposed list; recheck future course code table"},
        "worksheets": reports,
        "cross_sheet_invoice_multisets": {"fields": ["StockCode", "Description", "Quantity", "InvoiceDate", "Price", "Customer ID", "Country"], "shared_invoice_count": len(overlap_second), "identical_invoice_count": len(equal), "different_invoice_count": len(unequal), "removed_first_sheet_rows": removed, "removed_and_negative_regular": overlap_negative, "removed_and_nonproduct": overlap_nonproduct},
        "r1_for_existing_two_exclusion_rules": {"raw_rows": total, "cross_sheet_duplicate_rows": removed, "negative_regular_rows": negative, "overlap_rows": overlap_negative, "excluded_union_rows": excluded, "retained_rows": total - excluded, "note": "Preparation-rule count only; no processed data or database was produced."},
        "normalization": {"normalized_groups_with_multiple_raw_spellings": len(collisions), "examples": dict(list(collisions.items())[:12]), "note": "Spelling collisions alone do not establish business alias equivalence. No zero padding or customer merging performed."},
        "elapsed_seconds": round(time.perf_counter() - started, 3),
        "not_tested": ["matching rate", "price tolerance validity", "workflows", "concurrency", "10000-row end-to-end performance"],
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(args.output), "measured_at": result["measured_at"], "r1": result["r1_for_existing_two_exclusion_rules"], "identical_invoices": len(equal), "source_unchanged": True}, ensure_ascii=False))


if __name__ == "__main__":
    main()
