"""Read-only inspection of key collisions; no application or rule changes."""
import hashlib
import json
import unicodedata
import sys
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook

BASE = Path(__file__).resolve().parents[3]
SOURCE = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else BASE / 'VC-R-03-退货对账与差异处理系统/data/raw/online_retail_II.xlsx'
OUT = Path(sys.argv[2]).resolve() if len(sys.argv) > 2 else Path(__file__).with_name('key_collision_inspection.json')
digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
fields = {'Invoice': 0, 'StockCode': 1, 'Customer ID': 6}
groups = {name: defaultdict(set) for name in fields}
forbidden = {0x200B, 0x200C, 0x200D, 0x2060, 0xFEFF}
hits = {name: 0 for name in fields}
def narrow(value):
    width = value.translate({i: i - 0xFEE0 for i in range(0xFF01, 0xFF5F)} | {0x3000: 0x20})
    return width.strip().translate({i: i - 32 for i in range(0x61, 0x7B)})
wb = load_workbook(SOURCE, read_only=True, data_only=True)
rows = 0
for ws in wb:
    for values in ws.iter_rows(min_row=2, values_only=True):
        rows += 1
        for name, column in fields.items():
            raw = values[column]
            if raw is None:
                continue
            text = str(raw)
            key = unicodedata.normalize('NFKC', text).strip().upper()
            groups[name][key].add(text)
            hits[name] += any(ord(char) in forbidden for char in text)
wb.close()
output = {'measured_at': datetime.now(timezone(timedelta(hours=8))).isoformat(),
          'scope': 'Both complete raw worksheets before deduplication; counts are distinct normalized keys, not matched rows',
          'source_sha256': digest, 'source_rows': rows, 'unicode_version': unicodedata.unidata_version,
          'forbidden_codepoint_hits': hits, 'fields': {}}
for name, values in groups.items():
    merged = {key: spellings for key, spellings in values.items() if len(spellings) > 1}
    other = {key: spellings for key, spellings in merged.items() if len({narrow(s) for s in spellings}) > 1}
    output['fields'][name] = {'merged_keys': len(merged), 'explained_by_ascii_case_width_boundary_whitespace': len(merged) - len(other),
        'other_compatibility_collisions': len(other), 'other_examples': {k: sorted(v) for k,v in list(other.items())[:5]}}
output['source_unchanged'] = digest == hashlib.sha256(SOURCE.read_bytes()).hexdigest()
output['inspection_script_sha256'] = hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
assert output['source_unchanged']
OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps(output, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print(json.dumps(output, ensure_ascii=False, indent=2))
