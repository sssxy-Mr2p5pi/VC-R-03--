"""Dry-run the approved candidate rule against the immutable course workbook.

Produces design evidence only. No workflow database, allocation ledger, or
processed transaction file is created. A zero balance ledger is assumed.
"""
from __future__ import annotations

import argparse
import bisect
import csv
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from time import perf_counter

from audit_rule_inputs import EXACT
from openpyxl import load_workbook


def norm(value) -> str:
    return "" if value is None else unicodedata.normalize("NFKC", str(value)).strip().upper()


def nonproduct(row) -> bool:
    return norm(row[0]).startswith("A") or norm(row[1]) in EXACT or bool(re.fullmatch(r"TEST\d+|GIFT_\d+_\d+", norm(row[1])))


def checksum(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("source", type=Path)
    ap.add_argument("output_dir", type=Path)
    args = ap.parse_args()
    start = perf_counter()
    initial_hash = checksum(args.source)
    wb = load_workbook(args.source, read_only=True, data_only=True)
    sheets = wb.sheetnames
    raw = []
    invoice_rows = [defaultdict(list), defaultdict(list)]
    for si, ws in enumerate(wb.worksheets):
        it = ws.iter_rows(values_only=True)
        header = next(it)
        assert tuple(header) == ("Invoice", "StockCode", "Description", "Quantity", "InvoiceDate", "Price", "Customer ID", "Country")
        for rn, values in enumerate(it, 2):
            record = (si, rn, values)
            raw.append(record)
            invoice_rows[si][norm(values[0])].append(values)
    wb.close()
    shared = invoice_rows[0].keys() & invoice_rows[1].keys()
    equal = {inv for inv in shared if Counter(tuple(r[1:]) for r in invoice_rows[0][inv]) == Counter(tuple(r[1:]) for r in invoice_rows[1][inv])}
    assert len(shared - equal) == 0, "Different same-number invoice versions require a separate rule"
    imported = []
    excluded = Counter()
    for si, rn, row in raw:
        if si == 0 and norm(row[0]) in equal:
            excluded["CROSS_SHEET_DUPLICATE"] += 1
        elif not norm(row[0]).startswith("C") and row[3] < 0:
            excluded["NON_C_NEGATIVE_QUANTITY"] += 1
        else:
            imported.append((si, rn, row))
    raw_count = len(raw)
    del raw, invoice_rows
    index = defaultdict(list)
    originals = {}
    cancellation_rows = []
    invoices = defaultdict(lambda: {"commodity": [], "nonproduct_rows": 0})
    kept_nonproduct = 0
    for si, rn, row in imported:
        inv = norm(row[0])
        np = nonproduct(row)
        kept_nonproduct += int(np)
        if inv.startswith("C"):
            if np:
                invoices[inv]["nonproduct_rows"] += 1
            else:
                cancellation_rows.append((si, rn, row))
            continue
        if not np and row[3] > 0 and row[6] is not None:
            key = (norm(row[6]), norm(row[1]))
            original_id = (si, rn)
            originals[original_id] = row
            index[key].append((row[4], si, rn))
    for key in index:
        index[key].sort()
    dates = {key: [r[0] for r in values] for key, values in index.items()}
    results = []
    unique_claims = defaultdict(list)
    price_valid_claims = defaultdict(list)
    price_hist = Counter()
    source_hint_count = 0
    for si, rn, row in cancellation_rows:
        inv, sku, _desc, qty, when, price, customer, _country = row
        assert qty < 0 and price > 0 and isinstance(when, datetime), "Unexpected invalid cancellation input; extend audit explicitly"
        result = {"invoice": norm(inv), "sheet": sheets[si], "excel_row": rn, "raw_stock_code": str(sku), "normalized_stock_code": norm(sku), "raw_customer": customer, "normalized_customer": norm(customer), "candidate_count": 0, "reasons": [], "original": None, "price_delta": None, "requested_quantity": -qty}
        key = (norm(customer), norm(sku))
        candidates = index.get(key, [])
        end = bisect.bisect_right(dates.get(key, []), when)
        selected = candidates[:end] if customer is not None else []
        result["candidate_count"] = len(selected)
        if customer is None:
            result["reasons"].append("MISSING_CUSTOMER")
        elif not selected:
            result["reasons"].append("ONLY_LATER_SALES" if candidates else "NO_SAME_CUSTOMER_STOCK_IN_DATASET")
        elif len(selected) > 1:
            result["reasons"].append("MULTIPLE_CANDIDATES")
            quantities = [originals[(s, r)][3] for _, s, r in selected]
            if max(quantities) < -qty <= sum(quantities):
                result["reasons"].append("COMBINATION_POSSIBLE_NOT_PROVEN")
                source_hint_count += 1
        else:
            _, os, orn = selected[0]
            oid = (os, orn)
            original = originals[oid]
            delta = Decimal(str(price)) - Decimal(str(original[5]))
            result["original"] = {"invoice": norm(original[0]), "sheet": sheets[os], "excel_row": orn, "quantity": original[3], "price": str(original[5])}
            result["price_delta"] = str(delta)
            price_hist[str(delta)] += 1
            if price <= 0 or original[5] <= 0:
                result["reasons"].append("PRICE_NONPOSITIVE")
            elif delta != 0:
                result["reasons"].append("PRICE_DIFFERENCE")
            if -qty > original[3]:
                result["reasons"].append("SINGLE_RETURN_EXCEEDS_ORIGINAL")
            unique_claims[oid].append((len(results), -qty))
            if price > 0 and original[5] > 0 and delta == 0:
                price_valid_claims[oid].append((len(results), -qty))
        results.append(result)
    conflict_indices = set()
    conflicted_originals = []
    for oid, claims in unique_claims.items():
        requested = sum(qty for _, qty in claims)
        if requested > originals[oid][3]:
            conflicted_originals.append({"sheet": sheets[oid[0]], "excel_row": oid[1], "original_quantity": originals[oid][3], "total_unique_candidate_demand": requested, "cancellation_rows": len(claims)})
            conflict_indices.update(i for i, _ in claims)
    alternative_conflicts = {i for oid, claims in price_valid_claims.items() if sum(q for _, q in claims) > originals[oid][3] for i, _ in claims}
    for i in conflict_indices:
        results[i]["reasons"].append("CUMULATIVE_OVER_QUOTA")
    classes = Counter()
    primary_reasons = Counter()
    reason_tags = Counter()
    for row in results:
        row["classification"] = "NO_ORIGINAL" if row["candidate_count"] == 0 else "SUSPECTED" if row["reasons"] else "MATCHED_AUTO"
        row["primary_reason"] = row["reasons"][0] if row["reasons"] else "MATCHED_AUTO"
        classes[row["classification"]] += 1
        primary_reasons[row["primary_reason"]] += 1
        reason_tags.update(row["reasons"])
        invoices[row["invoice"]]["commodity"].append(row)
    invoice_classes = Counter()
    work_order_detail_rows = 0
    work_order_nonproduct_source_rows = 0
    pure_nonproduct_source_rows = 0
    automatic_invoice_product_rows = 0
    automatic_invoice_nonproduct_rows = 0
    for inv, content in invoices.items():
        statuses = {r["classification"] for r in content["commodity"]}
        classification = "NONPRODUCT_ONLY" if not statuses else "NO_ORIGINAL" if "NO_ORIGINAL" in statuses else "SUSPECTED" if "SUSPECTED" in statuses else "MATCHED_AUTO"
        invoice_classes[classification] += 1
        if classification in {"NO_ORIGINAL", "SUSPECTED"}:
            work_order_detail_rows += len(content["commodity"])
            work_order_nonproduct_source_rows += content["nonproduct_rows"]
        elif classification == "NONPRODUCT_ONLY":
            pure_nonproduct_source_rows += content["nonproduct_rows"]
        else:
            automatic_invoice_product_rows += len(content["commodity"])
            automatic_invoice_nonproduct_rows += content["nonproduct_rows"]
    # Independent linear filtering of the unsorted candidate sets for 32 spread-out inputs.
    sample_indices = sorted({int(i * (len(cancellation_rows) - 1) / 31) for i in range(32)})
    for i in sample_indices:
        _, _, row = cancellation_rows[i]
        expected = sum(1 for r in originals.values() if row[6] is not None and norm(r[6]) == norm(row[6]) and norm(r[1]) == norm(row[1]) and r[4] <= row[4])
        assert expected == results[i]["candidate_count"], (i, expected, results[i]["candidate_count"])
    automatic_allocations = Counter()
    for row in results:
        if row["classification"] == "MATCHED_AUTO":
            o = row["original"]
            automatic_allocations[(sheets.index(o["sheet"]), o["excel_row"])] += row["requested_quantity"]
    assert all(qty <= originals[oid][3] for oid, qty in automatic_allocations.items())
    assert raw_count == len(imported) + sum(excluded.values())
    assert sum(classes.values()) == len(results) == sum(primary_reasons.values())
    assert sum(invoice_classes.values()) == len(invoices)
    assert len(results) == work_order_detail_rows + automatic_invoice_product_rows
    cancellation_nonproduct_rows = sum(v['nonproduct_rows'] for v in invoices.values())
    assert cancellation_nonproduct_rows == work_order_nonproduct_source_rows + automatic_invoice_nonproduct_rows + pure_nonproduct_source_rows
    assert checksum(args.source) == initial_hash
    args.output_dir.mkdir(parents=True, exist_ok=True)
    csv_path = args.output_dir / "batch_line_results.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as out:
        writer = csv.DictWriter(out, fieldnames=list(results[0]))
        writer.writeheader()
        for row in results:
            writer.writerow({**row, "reasons": json.dumps(row["reasons"], ensure_ascii=False), "original": json.dumps(row["original"], ensure_ascii=False)})
    examples = {}
    for row in results:
        for reason in row["reasons"] or ["MATCHED_AUTO"]:
            examples.setdefault(reason, row)
    summary = {
        "measured_at": datetime.now(timezone(timedelta(hours=8))).isoformat(),
        "source_sha256": initial_hash, "source_unchanged": True,
        "scope": "Both sheets, identical invoice multisets deduplicated retaining second sheet, non-C negative quantities excluded. Full historical candidate pool, empty effective ledger; design dry-run only.",
        "rule": {"normalization": "NFKC + strip ends + uppercase; preserve original; no aliases or zero padding", "candidate": "same customer + stock, positive original quantity, original time <= cancellation time", "price": "Decimal source unit price, GBP, tolerance zero; no tax assumption", "quota": "Sum ALL uniquely referenced candidate quantities including price-mismatch rows; mark all references on over-quota; no temporal priority", "multiple_candidates": "No allocation", "combination": "Hint only; no allocation", "nonproduct_dictionary": "Same exact code list as audit_rule_inputs.py; course update pending"},
        "row_counts": {"raw": raw_count, "imported": len(imported), "excluded": dict(excluded), "nonproduct_retained_both_sides": kept_nonproduct, "commodity_cancellation": len(results), "commodity_classes": dict(classes), "primary_reasons_exclusive": dict(primary_reasons), "reason_tags_nonexclusive": dict(reason_tags)},
        "invoice_counts": dict(invoice_classes),
        "work_order_scope": {"predicted_work_orders": invoice_classes["NO_ORIGINAL"] + invoice_classes["SUSPECTED"], "all_commodity_lines_in_predicted_work_orders": work_order_detail_rows, "nonproduct_source_rows_in_work_order_footers": work_order_nonproduct_source_rows, "pure_nonproduct_source_rows": pure_nonproduct_source_rows, "automatic_invoice_product_rows": automatic_invoice_product_rows, "automatic_invoice_nonproduct_rows": automatic_invoice_nonproduct_rows, "all_cancellation_nonproduct_rows": cancellation_nonproduct_rows, "note": "Predictions only; no tickets created. Commodity lines of fully automatic invoices remain in results, not work-order details."},
        "quota": {"conflicted_original_rows": len(conflicted_originals), "conflicted_cancellation_rows": len(conflict_indices), "combination_hint_rows": source_hint_count, "conflict_examples": conflicted_originals[:5], "alternative_price_valid_demand_only": {"conflicted_cancellation_rows": len(alternative_conflicts), "different_rows": len(conflict_indices ^ alternative_conflicts), "note": "Sensitivity comparison only; this alternate rule is not adopted."}},
        "price_differences_unique_candidates": {"unique_rows": sum(price_hist.values()), "exactly_equal": sum(v for k, v in price_hist.items() if Decimal(k) == 0), "nonzero_abs_at_most_0_01": sum(v for k, v in price_hist.items() if 0 < abs(Decimal(k)) <= Decimal('0.01')), "greater_than_0_01": sum(v for k, v in price_hist.items() if abs(Decimal(k)) > Decimal('0.01')), "note": "No ground truth correspondence; price distribution cannot establish a legitimate tolerance."},
        "checks": {"source_hash_unchanged": True, "row_balance": True, "classification_balance": True, "invoice_balance": True, "linear_candidate_crosschecks": len(sample_indices), "proposed_automatic_allocations_within_original_quantity": True},
        "examples": examples,
        "artifacts": {"line_results": csv_path.name, "line_results_sha256": checksum(csv_path), "audit_script_sha256": checksum(Path(__file__)), "classifier_script_sha256": checksum(Path(__file__).with_name('audit_rule_inputs.py'))},
        "elapsed_seconds": round(perf_counter() - start, 3),
        "limitations": ["No workflow/ledger transactions were executed", "No application UI or 10000-row performance acceptance was tested", "No matching correctness ground truth exists", "Nonproduct dictionary remains subject to the future course update"],
    }
    output = args.output_dir / "batch_validation.json"
    output.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"summary": str(output), "rows": summary["row_counts"], "invoices": summary["invoice_counts"], "quota": summary["quota"], "price": summary["price_differences_unique_candidates"], "checks": summary["checks"], "elapsed_seconds": summary["elapsed_seconds"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
