"""Read every original worksheet without altering, filtering, or pairing rows."""
from __future__ import annotations

import csv
import json
import time
from collections import Counter
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook

from download_data import ROOT, WORKBOOK_NAME, sha256

ALIASES = {
    "invoice": ("Invoice", "InvoiceNo"),
    "stock_code": ("StockCode",),
    "description": ("Description",),
    "quantity": ("Quantity",),
    "invoice_date": ("InvoiceDate",),
    "unit_price": ("Price", "UnitPrice"),
    "customer_id": ("Customer ID", "CustomerID"),
    "country": ("Country",),
}


def empty(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def serial(value: object) -> object:
    return value.isoformat() if isinstance(value, (datetime, date)) else value


def profile(path: Path) -> dict:
    started = time.perf_counter()
    before = sha256(path)
    source_manifest = json.loads((ROOT / "evidence/data/download_manifest.json").read_text(encoding="utf-8"))
    expected_hash = next(item["sha256"] for item in source_manifest["files"] if item["path"].endswith(".xlsx"))
    if before != expected_hash:
        raise ValueError("Raw workbook hash differs from the download manifest.")
    sheets = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    preview = []
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = list(next(rows))
            indices = {}
            actual_mapping = {}
            for target, names in ALIASES.items():
                found = [(i, name) for i, name in enumerate(header) if name in names]
                if len(found) != 1:
                    raise ValueError(f"{sheet.title}: missing or ambiguous column for {target}: {header}")
                indices[target], actual_mapping[target] = found[0]
            counts = Counter()
            missing = Counter()
            month_rows = Counter()
            first = last = None
            date_min = date_max = None
            for row_number, row in enumerate(rows, 2):
                counts["physical_rows_after_header"] += 1
                if all(empty(cell) for cell in row):
                    counts["fully_blank_rows"] += 1
                    continue
                counts["data_rows"] += 1
                record = {name: row[index] for name, index in indices.items()}
                for name, value in record.items():
                    if empty(value):
                        missing[name] += 1
                locator = {"sheet": sheet.title, "excel_row": row_number, "values": [serial(value) for value in row]}
                first = first or locator
                last = locator
                if counts["data_rows"] <= 5:
                    preview.append([sheet.title, row_number] + [serial(value) for value in row])
                invoice = record["invoice"]
                if empty(invoice):
                    counts["invoice_missing"] += 1
                    is_cancel = None
                else:
                    is_cancel = str(invoice).upper().startswith("C")
                    counts["c_prefix_rows" if is_cancel else "non_c_prefix_rows"] += 1
                quantity = record["quantity"]
                if isinstance(quantity, (int, float)) and not isinstance(quantity, bool):
                    counts["quantity_negative" if quantity < 0 else "quantity_zero" if quantity == 0 else "quantity_positive"] += 1
                    if is_cancel is False and quantity < 0:
                        counts["non_c_with_negative_quantity"] += 1
                    if is_cancel is True and quantity >= 0:
                        counts["c_with_nonnegative_quantity"] += 1
                else:
                    counts["quantity_missing_or_non_numeric"] += 1
                price = record["unit_price"]
                if isinstance(price, (int, float)) and not isinstance(price, bool):
                    counts["price_negative" if price < 0 else "price_zero" if price == 0 else "price_positive"] += 1
                else:
                    counts["price_missing_or_non_numeric"] += 1
                when = record["invoice_date"]
                if isinstance(when, datetime):
                    date_min = when if date_min is None else min(date_min, when)
                    date_max = when if date_max is None else max(date_max, when)
                    month_rows[when.strftime("%Y-%m")] += 1
                else:
                    counts["date_missing_or_non_datetime"] += 1
            if counts["data_rows"] != counts["c_prefix_rows"] + counts["non_c_prefix_rows"] + counts["invoice_missing"]:
                raise ValueError("Invoice classification row accounting failed.")
            sheets.append({
                "sheet": sheet.title,
                "header_excel_row": 1,
                "actual_headers": header,
                "field_mapping": actual_mapping,
                "counts": dict(counts),
                "missing_by_field": {name: missing[name] for name in ALIASES},
                "date_min": serial(date_min), "date_max": serial(date_max),
                "rows_by_month": dict(sorted(month_rows.items())),
                "first_data_row": first, "last_data_row": last,
            })
            print(f"{sheet.title}: {counts['data_rows']:,} rows read", flush=True)
    finally:
        workbook.close()
    after = sha256(path)
    if before != after:
        raise ValueError("Original workbook changed during profiling.")
    total = sum(sheet["counts"]["data_rows"] for sheet in sheets)
    expected = source_manifest["official_expected_total_rows"]
    result = {
        "profiled_at": datetime.now(timezone(timedelta(hours=8))).isoformat(timespec="seconds"),
        "source_path": str(path.relative_to(ROOT)).replace("\\", "/"),
        "source_sha256_before": before, "source_sha256_after": after,
        "header_policy": "Excel row 1 explicitly treated as the header; every remaining row inspected in every sheet.",
        "total_data_rows": total, "official_expected_total_rows": expected,
        "matches_official_total": total == expected,
        "sheets": sheets,
        "duration_seconds": round(time.perf_counter() - started, 3),
        "scope": "Read-only preparation profile. No cleaning, slicing, reconciliation, assignments, or business acceptance performed.",
        "not_measured": ["duplicate row count", "ground-truth return links", "business end-to-end performance"],
    }
    destination = ROOT / "evidence/data"
    destination.mkdir(parents=True, exist_ok=True)
    (destination / "raw_profile.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    with (destination / "raw_preview.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["source_sheet", "source_excel_row"] + sheets[0]["actual_headers"])
        writer.writerows(preview)
    if total != expected:
        raise ValueError(f"Actual {total:,} rows differ from official expected {expected:,}; inspect the saved profile.")
    return result


def main() -> None:
    result = profile(ROOT / "data/raw" / WORKBOOK_NAME)
    print(json.dumps({"total_data_rows": result["total_data_rows"], "matches_official_total": result["matches_official_total"], "read_only_duration_seconds": result["duration_seconds"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
