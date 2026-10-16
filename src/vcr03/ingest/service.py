"""One atomic import service for CLI and Flask; no matching or workflow."""
from __future__ import annotations

from collections import Counter, defaultdict, deque
from datetime import datetime, timedelta, timezone
from decimal import Decimal, Inexact, Rounded, localcontext
import hashlib
import importlib.metadata
import json
from pathlib import Path
import sqlite3
from time import perf_counter
import unicodedata
from uuid import uuid4

from openpyxl import load_workbook

from vcr03.ingest.dictionary import Dictionary
from vcr03.ingest.rules import Cell, HEADERS, INGEST_VERSION, NORM_VERSION, READ_VERSION, compact, digest, packed, parse, read_keys, route, typed_content, unpack_cells
from vcr03.storage.database import connect


class ImportRejected(ValueError):
    """Input or identity requires correction/approval before activation."""


def file_hash(path: Path):
    h = hashlib.sha256()
    with path.open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def now():
    return datetime.now(timezone(timedelta(hours=8))).isoformat(timespec='seconds')


def ab_binding_conflict(db, dataset_id, snapshot_id):
    """Return the A snapshot id bound to the current ACTIVE B when it differs, else None.

    Pure-A databases have no matching tables (treated as no binding, never an error).
    A binding whose matching_snapshot is not ACTIVE (e.g. REQUIRES_REVIEW) or a missing
    active_matching row never conflicts; the same snapshot id passes (idempotent reuse).
    """
    tables = {row[0] for row in db.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND name IN ('active_matching','matching_snapshot')")}
    if not {'active_matching', 'matching_snapshot'} <= tables:
        return None
    bound = db.execute('''SELECT s.snapshot_id,s.status FROM active_matching m
                          JOIN matching_snapshot s ON s.match_id=m.match_id
                          WHERE m.dataset_id=?''', (dataset_id,)).fetchone()
    if not bound or bound['status'] != 'ACTIVE' or bound['snapshot_id'] == snapshot_id:
        return None
    return bound['snapshot_id']


ANALYSIS_INSERT = '''INSERT INTO line_analysis VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)'''


def analysis_values(snapshot, line_id, row):
    return (snapshot, line_id, row['keys']['Invoice'], row['keys']['StockCode'], row['keys']['Customer ID'],
            NORM_VERSION, row['side'], row['line_kind'], row['quantity'], row['price'], row['invoice_date'],
            'RETAINED', None, '[]', None, row['reason_category'], row['primary_reason'], compact(row['issues']),
            row['candidate_comparable'])


def ingest_workbook(source: Path, database: Path, *, dictionary_path=None, progress=None,
                    approved_unchanged_rule_refresh=False):
    """Import the entire two-sheet XLSX, or reuse an identical successful snapshot."""
    started = perf_counter()
    source, database = Path(source).resolve(), Path(database).resolve()
    if unicodedata.unidata_version != '15.0.0':
        raise ImportRejected('运行环境须使用Python 3.12 / Unicode 15.0.0。')
    if importlib.metadata.version('openpyxl') != '3.1.5':
        raise ImportRejected('读取合同固定openpyxl 3.1.5，请按锁定依赖重建环境。')
    if source == database or source.suffix.lower() != '.xlsx':
        raise ImportRejected('输入须为独立XLSX，不能与数据库重合。')
    initial_hash = file_hash(source)
    dictionary = Dictionary(dictionary_path)
    rules = {'dictionary': dictionary.content, 'dictionary_hash': dictionary.hash,
             'normalization': NORM_VERSION, 'reading': READ_VERSION, 'ingestion': INGEST_VERSION,
             'deduplication': 'typed-seven-approved-identifiers-v1'}
    rule_hash = digest(rules)
    workbook = load_workbook(source, read_only=True, data_only=False)
    db = None
    try:
        if len(workbook.worksheets) != 2:
            raise ImportRejected('首版输入须有两张同结构数据表。')
        ranges = []
        for ws in workbook.worksheets:
            header = tuple(c.value for c in next(ws.iter_rows(min_row=1, max_row=1, max_col=8)))
            if header != HEADERS or ws.max_column != 8:
                raise ImportRejected(f'{ws.title}表头与固定八列不一致。')
            ranges.append({'sheet': ws.title, 'first_row': 2, 'last_row': ws.max_row})
        dataset = digest({'file': initial_hash, 'ranges': ranges})
        snapshot = digest({'dataset': dataset, 'rules': rule_hash})
        db = connect(database, initialize=True)
        db.execute('BEGIN IMMEDIATE')
        existing = db.execute('SELECT source_hash FROM source_dataset').fetchall()
        if any(r['source_hash'] != initial_hash for r in existing):
            raise ImportRejected('SOURCE_VERSION_REQUIRES_APPROVAL：不同源哈希不能自动并入当前库，请使用独立库。')
        found = db.execute('SELECT * FROM analysis_snapshot WHERE snapshot_id=?', (snapshot,)).fetchone()
        if found:
            if found['status'] != 'A_COMPLETE':
                raise ImportRejected('已有分析未完整完成，不能复用。')
            active = db.execute('SELECT snapshot_id FROM active_analysis WHERE dataset_id=?', (dataset,)).fetchone()
            if active and active['snapshot_id'] != snapshot:
                raise ImportRejected('RULE_CHANGE_REQUIRES_APPROVAL：不能自动激活历史规则快照。')
            bound = ab_binding_conflict(db, dataset, snapshot)
            if bound:
                raise ImportRejected('A_BINDING_CONFLICT_REQUIRES_APPROVAL：'
                                     f'本次待激活A快照{snapshot}与当前有效B绑定的A快照{bound}不一致，'
                                     '本次激活已拒绝；既有来源、分析、配对、额度与工单保持。')
            db.execute('INSERT OR IGNORE INTO active_analysis VALUES(?,?)', (dataset, snapshot))
            summary = json.loads(found['summary_json'])
            db.execute('INSERT INTO import_run VALUES(?,?,?,?,?)', (uuid4().hex, snapshot, now(), 1, perf_counter()-started))
            db.commit()
            return {'dataset_id': dataset, 'snapshot_id': snapshot, 'reused': True,
                    'elapsed_seconds': perf_counter()-started, 'summary': summary}
        previous = db.execute("""SELECT a.snapshot_id FROM analysis_snapshot a
                                 LEFT JOIN active_analysis x ON x.dataset_id=a.dataset_id
                                 WHERE a.dataset_id=? AND a.status='A_COMPLETE'
                                 ORDER BY (a.snapshot_id=x.snapshot_id) DESC,a.created_at DESC LIMIT 1""", (dataset,)).fetchone()
        if previous and not approved_unchanged_rule_refresh:
            raise ImportRejected('RULE_CHANGE_REQUIRES_APPROVAL：新规则须先比较确认，不能自动覆盖激活。')
        if not previous:
            db.execute('INSERT INTO source_dataset VALUES(?,?,?,?,?,?)',
                       (dataset, initial_hash, source.name, compact(ranges), workbook.epoch.isoformat(), now()))
        db.execute('INSERT INTO analysis_snapshot VALUES(?,?,?,?,?,?,?,?,?,?,?)',
                   (snapshot, dataset, rule_hash, dictionary.version, dictionary.hash, NORM_VERSION,
                    READ_VERSION, compact(rules), 'IMPORTING', '{}', now()))
        variants = {}
        source_batch, analysis_batch = [], []
        line_id = 0
        # Phase 1: immutable raw values/types and all three keys. No business-value validation yet.
        for ws in workbook.worksheets:
            for rn, cells in enumerate(ws.iter_rows(min_row=2, max_col=8), 2):
                line_id += 1
                input_cells = tuple(Cell(c.value, c.data_type) for c in cells)
                keys, narrows, issues = read_keys(input_cells)
                row = {'keys': keys, 'issues': issues, 'quantity': None, 'price': None, 'invoice_date': None,
                       'side': 'UNVALIDATED', 'line_kind': 'UNVALIDATED', 'reason_category': None,
                       'primary_reason': None, 'candidate_comparable': 0}
                raw_cells = [{'type': c.kind, 'value': packed(c.value)} for c in input_cells]
                if not previous:
                    source_batch.append((line_id, dataset, ws.title, rn, compact(raw_cells)))
                analysis_batch.append(analysis_values(snapshot, line_id, row))
                for field, pos in [('Invoice', 0), ('StockCode', 1), ('Customer ID', 6)]:
                    k = keys[field]
                    if k:
                        variant = (field, k, compact(raw_cells[pos]), narrows[field])
                        if variant not in variants:
                            variants[variant] = [0, line_id]
                        variants[variant][0] += 1
                if len(analysis_batch) >= 1000:
                    if source_batch:
                        db.executemany('INSERT INTO source_line VALUES(?,?,?,?,?)', source_batch)
                    db.executemany(ANALYSIS_INSERT, analysis_batch)
                    source_batch.clear()
                    analysis_batch.clear()
                if progress and line_id % 100000 == 0:
                    progress({'phase': 'read-keys', 'rows': line_id})
        if analysis_batch:
            if source_batch:
                db.executemany('INSERT INTO source_line VALUES(?,?,?,?,?)', source_batch)
            db.executemany(ANALYSIS_INSERT, analysis_batch)
        expected = sum(max(0, r['last_row']-1) for r in ranges)
        if line_id != expected:
            raise ImportRejected('物理来源行数与工作表范围不一致。')
        narrow_groups = defaultdict(set)
        raw_groups = defaultdict(set)
        for field, k, raw, narrow in variants:
            narrow_groups[(field, k)].add(narrow)
            raw_groups[(field, k)].add(raw)
        suspicious = {g for g, values in narrow_groups.items() if len(values) > 1}
        db.executemany('INSERT INTO key_variant VALUES(?,?,?,?,?,?,?,?)',
                       [(snapshot, field, k, raw, narrow, count, first, int((field, k) in suspicious))
                        for (field, k, raw, narrow), (count, first) in variants.items()])
        # Phase 2: full multisets; invalid business values are still compared with their types.
        groups = [defaultdict(lambda: defaultdict(list)) for _ in range(2)]
        type_counts = [defaultdict(Counter) for _ in range(2)]
        uncomparable = set()
        sheet_indexes = {r['sheet']: i for i, r in enumerate(ranges)}
        select = """SELECT s.line_id,s.sheet,s.raw_json,a.invoice_key,a.stock_key,a.customer_key
                    FROM source_line s JOIN line_analysis a USING(line_id)
                    WHERE s.dataset_id=? AND a.snapshot_id=? ORDER BY s.line_id"""
        for stored in db.execute(select, (dataset, snapshot)):
            invoice = stored['invoice_key']
            if not invoice or ('Invoice', invoice) in suspicious:
                continue
            cells = unpack_cells(json.loads(stored['raw_json']))
            keys = {'Invoice': invoice, 'StockCode': stored['stock_key'], 'Customer ID': stored['customer_key']}
            content = typed_content(cells, keys, suspicious, normalized_identifiers=True)
            si = sheet_indexes[stored['sheet']]
            groups[si][invoice][content].append(stored['line_id'])
            type_counts[si][invoice][tuple(c.kind for c in cells[1:])] += 1
            if content is None:
                uncomparable.add(invoice)
        duplicates, conflicts = [], {}
        for invoice in sorted(groups[0].keys() & groups[1].keys()):
            earlier, later = groups[0][invoice], groups[1][invoice]
            if invoice not in uncomparable and Counter({r: len(ids) for r, ids in earlier.items()}) == Counter({r: len(ids) for r, ids in later.items()}):
                for content, ids in earlier.items():
                    keep = deque(later[content])
                    duplicates.extend((keep.popleft(), snapshot, old) for old in ids)
            else:
                conflicts[invoice] = ('INVOICE_DEDUP_TYPE_CONFLICT' if invoice in uncomparable or
                                      type_counts[0][invoice] != type_counts[1][invoice] else 'INVOICE_VERSION_CONFLICT')
        del groups, type_counts
        db.executemany("""UPDATE line_analysis SET disposition='EXCLUDED',exclusion_primary='CROSS_SHEET_DUPLICATE',
                          exclusion_tags='["CROSS_SHEET_DUPLICATE"]',retained_line_id=?
                          WHERE snapshot_id=? AND line_id=?""", duplicates)
        # Phase 3: value validation/routing follows the physical whole-invoice comparison.
        updates = []
        for stored in db.execute(select, (dataset, snapshot)):
            row = parse(unpack_cells(json.loads(stored['raw_json'])), dictionary)
            key_issues = [{'field': f, 'code': 'KEY_NORMALIZATION_CONFLICT'}
                          for f in ('Invoice', 'StockCode', 'Customer ID')
                          if (f, row['keys'][f]) in suspicious]
            if key_issues:
                row['issues'].extend(key_issues)
                row = route(row, dictionary)
            conflict = conflicts.get(row['keys']['Invoice'])
            if conflict:
                row['issues'].append({'field': 'Invoice', 'code': conflict})
                # D-38: version conflict blocks allocation, not candidate discovery.
                row['reason_category'] = ('NONPRODUCT_DATA_QUALITY' if row['line_kind'] == 'NONPRODUCT' else 'DATA_QUALITY')
                row['primary_reason'] = row['primary_reason'] or conflict
            updates.append((row['side'], row['line_kind'], row['quantity'], row['price'], row['invoice_date'],
                            row['reason_category'], row['primary_reason'], compact(row['issues']),
                            row['candidate_comparable'], snapshot, stored['line_id']))
            if len(updates) >= 1000:
                db.executemany("""UPDATE line_analysis SET side=?,line_kind=?,quantity=?,price=?,invoice_date=?,
                                  reason_category=?,primary_reason=?,quality_tags=?,candidate_comparable=?
                                  WHERE snapshot_id=? AND line_id=?""", updates)
                updates.clear()
        if updates:
            db.executemany("""UPDATE line_analysis SET side=?,line_kind=?,quantity=?,price=?,invoice_date=?,
                              reason_category=?,primary_reason=?,quality_tags=?,candidate_comparable=?
                              WHERE snapshot_id=? AND line_id=?""", updates)
        negative = db.execute("SELECT line_id,exclusion_primary FROM line_analysis WHERE snapshot_id=? AND side='SALE' AND quantity<0", (snapshot,)).fetchall()
        overlap = sum(r['exclusion_primary'] is not None for r in negative)
        db.executemany('''UPDATE line_analysis SET disposition='EXCLUDED',
                          exclusion_primary=COALESCE(exclusion_primary,'NON_C_NEGATIVE_QUANTITY'),
                          exclusion_tags=? WHERE snapshot_id=? AND line_id=?''',
                       [(compact(['CROSS_SHEET_DUPLICATE', 'NON_C_NEGATIVE_QUANTITY'] if r['exclusion_primary'] else
                                 ['NON_C_NEGATIVE_QUANTITY']), snapshot, r['line_id']) for r in negative])
        excluded = db.execute("SELECT COUNT(*) FROM line_analysis WHERE snapshot_id=? AND disposition='EXCLUDED'", (snapshot,)).fetchone()[0]
        quality = db.execute("SELECT COUNT(*) FROM line_analysis WHERE snapshot_id=? AND disposition='RETAINED' AND reason_category IS NOT NULL", (snapshot,)).fetchone()[0]
        side_counts = dict(db.execute("SELECT side,COUNT(*) FROM line_analysis WHERE snapshot_id=? AND disposition='RETAINED' GROUP BY side", (snapshot,)).fetchall())
        exclusion_counts = dict(db.execute("SELECT exclusion_primary,COUNT(*) FROM line_analysis WHERE snapshot_id=? AND disposition='EXCLUDED' GROUP BY exclusion_primary", (snapshot,)).fetchall())
        np_stats = {}
        with localcontext() as ctx:
            ctx.prec = 1000
            ctx.traps[Inexact] = True
            ctx.traps[Rounded] = True
            for r in db.execute("SELECT invoice_key,stock_key,side,quantity,price FROM line_analysis WHERE snapshot_id=? AND disposition='RETAINED' AND line_kind='NONPRODUCT'", (snapshot,)):
                code = r['stock_key'] or '(缺失)'
                item = np_stats.setdefault(code, {'code': code, **dictionary.description(r['invoice_key'], r['stock_key']),
                                                 'count': 0, 'sale': 0, 'cancel': 0, 'amount': Decimal(0), 'invalid_amount_rows': 0})
                item['count'] += 1
                item['sale' if r['side'] == 'SALE' else 'cancel'] += 1
                if r['quantity'] is None or r['price'] is None:
                    item['invalid_amount_rows'] += 1
                else:
                    item['amount'] += Decimal(r['quantity']) * Decimal(r['price'])
        for item in np_stats.values():
            item['amount'] = str(item['amount'])
        retained = db.execute("SELECT COUNT(*) FROM line_analysis WHERE snapshot_id=? AND disposition='RETAINED'", (snapshot,)).fetchone()[0]
        stored_count = db.execute('SELECT COUNT(*) FROM source_line WHERE dataset_id=?', (dataset,)).fetchone()[0]
        if line_id != stored_count or line_id != retained+excluded or excluded != len(duplicates)+len(negative)-overlap:
            raise ImportRejected('R1或排除并集验算失败。')
        if file_hash(source) != initial_hash:
            raise ImportRejected('SOURCE_CHANGED：导入时源文件发生变化，已回滚。')
        summary = {'raw_rows': line_id, 'retained_rows': retained, 'excluded_rows': excluded,
                   'r1_valid': True, 'source_sha256': initial_hash, 'source_unchanged': True,
                   'side_counts': side_counts, 'exclusion_primary_counts': exclusion_counts,
                   'exclusion_union': {'duplicate': len(duplicates), 'negative': len(negative), 'overlap': overlap},
                   'quality_rows': quality, 'nonproduct_rows': sum(i['count'] for i in np_stats.values()),
                   'nonproduct': sorted(np_stats.values(), key=lambda r: r['code']),
                   'key_merge_groups': sum(len(v) > 1 for v in raw_groups.values()),
                   'suspicious_key_groups': len(suspicious), 'invoice_version_conflicts': sorted(conflicts),
                   'invoice_type_conflicts': sorted(k for k, v in conflicts.items() if v == 'INVOICE_DEDUP_TYPE_CONFLICT'),
                   'dictionary_version': dictionary.version, 'dictionary_hash': dictionary.hash,
                   'normalization_version': NORM_VERSION, 'read_version': READ_VERSION,
                   'scope': 'module A import and quality only; no matching, ticket, allocation or NFR acceptance'}
        if previous:
            columns = [r['name'] for r in db.execute('PRAGMA table_info(line_analysis)')
                       if r['name'] not in {'snapshot_id', 'normalization_version'}]
            difference = ' OR '.join(f'n.{c} IS NOT o.{c}' for c in columns)
            changed = db.execute(f"""SELECT COUNT(*) FROM line_analysis n JOIN line_analysis o USING(line_id)
                                    WHERE n.snapshot_id=? AND o.snapshot_id=? AND ({difference})""",
                                 (snapshot, previous['snapshot_id'])).fetchone()[0]
            old_count = db.execute('SELECT COUNT(*) FROM line_analysis WHERE snapshot_id=?', (previous['snapshot_id'],)).fetchone()[0]
            if changed or old_count != line_id:
                raise ImportRejected(f'RULE_RESULTS_CHANGED_REQUIRES_CONFIRMATION：{changed}行输出变化，新旧行数{line_id}/{old_count}；新分析已回滚。')
            summary['approved_refresh'] = {'previous_snapshot_id': previous['snapshot_id'], 'changed_output_rows': 0,
                                           'compared_rows': line_id, 'old_source_and_analysis_preserved': True}
        bound = ab_binding_conflict(db, dataset, snapshot)
        if bound:
            raise ImportRejected('A_BINDING_CONFLICT_REQUIRES_APPROVAL：'
                                 f'本次待激活A快照{snapshot}与当前有效B绑定的A快照{bound}不一致，'
                                 '本次激活已拒绝；既有来源、分析、配对、额度与工单保持。')
        db.execute('INSERT INTO active_analysis VALUES(?,?) ON CONFLICT(dataset_id) DO UPDATE SET snapshot_id=excluded.snapshot_id',
                   (dataset, snapshot))
        db.execute("UPDATE analysis_snapshot SET status='A_COMPLETE',summary_json=? WHERE snapshot_id=?", (compact(summary), snapshot))
        db.execute('INSERT INTO import_run VALUES(?,?,?,?,?)', (uuid4().hex, snapshot, now(), 0, perf_counter()-started))
        db.commit()
        saved = json.loads(db.execute('SELECT summary_json FROM analysis_snapshot WHERE snapshot_id=?', (snapshot,)).fetchone()[0])
        return {'dataset_id': dataset, 'snapshot_id': snapshot, 'reused': False,
                'elapsed_seconds': perf_counter()-started, 'summary': saved}
    except sqlite3.OperationalError as exc:
        if db:
            db.rollback()
        if 'locked' in str(exc).lower():
            raise ImportRejected('DATABASE_BUSY：导入写入正在进行，请稍后重试。') from exc
        raise
    except Exception:
        if db:
            db.rollback()
        raise
    finally:
        workbook.close()
        if db:
            db.close()
