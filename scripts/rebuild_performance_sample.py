"""D-37: recreate the existing fixed sample, preserving its byte identity."""
from datetime import datetime
from decimal import Decimal
from pathlib import Path
import csv
import hashlib
import json
import sys
from uuid import uuid4
from zipfile import ZipFile, ZipInfo

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'src'))
from vcr03.ingest.service import file_hash  # noqa: E402
from vcr03.ingest.rules import HEADERS  # noqa: E402


def capture_recipe(sample):
    """Read-only packaging inventory; original dates recreate bytes, never date new work."""
    with ZipFile(sample) as z:
        return {'captured_at':datetime.now().astimezone().isoformat(),'sample_sha256':file_hash(sample),
                'purpose':'Original packaging only; rebuild execution time is recorded separately',
                'core_xml':z.read('docProps/core.xml').decode('utf-8'),'archive_comment':z.comment.hex(),
                'members':[{'name':i.filename,'date_time':list(i.date_time),'compress_type':i.compress_type,
                            'external_attr':i.external_attr,'create_system':i.create_system,
                            'create_version':i.create_version,'extract_version':i.extract_version,
                            'flag_bits':i.flag_bits,'internal_attr':i.internal_attr,
                            'extra':i.extra.hex(),'comment':i.comment.hex(),
                            'sha256':hashlib.sha256(z.read(i)).hexdigest()}
                           for i in z.infolist()]}


def restore_archive(generated,target,recipe):
    """Verify every XML payload before writing, then verify the whole archive identity."""
    if target.exists():
        raise ValueError('已有样本禁止覆盖。')
    payloads=[]
    with ZipFile(generated) as z:
        if z.namelist()!=[i['name'] for i in recipe['members']]:
            raise ValueError('封装成员不一致，不能发布。')
        for member in recipe['members']:
            payload=(recipe['core_xml'].encode('utf-8') if member['name']=='docProps/core.xml'
                     else z.read(member['name']))
            if hashlib.sha256(payload).hexdigest()!=member['sha256']:
                raise ValueError(f"生成内容哈希不一致：{member['name']}；不能发布。")
            payloads.append((member,payload))
    with ZipFile(target,'x') as z:
        z.comment=bytes.fromhex(recipe['archive_comment'])
        for member,payload in payloads:
            info=ZipInfo(member['name'],tuple(member['date_time']))
            for name in ['compress_type','external_attr','create_system','create_version',
                         'extract_version','flag_bits','internal_attr']:
                setattr(info,name,member[name])
            info.extra=bytes.fromhex(member['extra'])
            info.comment=bytes.fromhex(member['comment'])
            z.writestr(info,payload,compresslevel=6)
    if file_hash(target)!=recipe['sample_sha256']:
        target.unlink()  # owned file created exclusively by this call; existing files never opened for write
        raise ValueError('既有样本字节哈希无法复现，停止发布。')


def rebuild_sample(source,output,evidence):
    source,output,evidence=Path(source),Path(output),Path(evidence)
    if output.exists():
        raise ValueError('已有样本禁止覆盖。')
    manifest_file=evidence/'manifest.json'
    mapping_file=evidence/'source_rows.csv'
    recipe_file=evidence/'rebuild-metadata.json'
    if not all(p.is_file() for p in [manifest_file,mapping_file,recipe_file]):
        raise ValueError('既有清单/映射/封装元数据缺失；停止，不重新选样。')
    preserved={p.name:file_hash(p) for p in [manifest_file,mapping_file,recipe_file]}
    manifest=json.loads(manifest_file.read_text(encoding='utf-8'))
    recipe=json.loads(recipe_file.read_text(encoding='utf-8'))
    if file_hash(source)!=manifest['source_sha256']:
        raise ValueError('源哈希不一致，停止重建。')
    if file_hash(mapping_file)!=manifest['mapping_sha256']:
        raise ValueError('来源映射哈希不一致，停止重建。')
    if recipe['sample_sha256']!=manifest['sample_sha256']:
        raise ValueError('封装元数据与原样本哈希不一致。')
    with mapping_file.open(encoding='utf-8-sig',newline='') as f:
        mapping=list(csv.DictReader(f))
    if len(mapping)!=manifest['rows'] or len(mapping)!=10000:
        raise ValueError('映射必须是既有10000行，不自动补齐。')
    selected={}
    wb=load_workbook(source,read_only=True,data_only=False)
    try:
        if [w.title for w in wb.worksheets]!=[w['sheet'] for w in manifest['worksheets']]:
            raise ValueError('来源工作表不一致。')
        epoch=wb.epoch
        for ws in wb.worksheets:
            entries=[r for r in mapping if r['source_sheet']==ws.title]
            original_rows=[int(r['source_excel_row']) for r in entries]
            if (len(entries)!=5000 or original_rows!=sorted(set(original_rows)) or
                [int(r['sample_excel_row']) for r in entries]!=list(range(2,5002)) or
                any(r['source_sha256']!=manifest['source_sha256'] or r['sample_sheet']!=ws.title for r in entries)):
                raise ValueError('映射身份/顺序/数量不一致。')
            if tuple(c.value for c in next(ws.iter_rows(max_row=1,max_col=8)))!=HEADERS:
                raise ValueError('来源表头不一致。')
            wanted=set(original_rows)
            chosen=[]
            for rn,cells in enumerate(ws.iter_rows(min_row=2,max_row=max(wanted),max_col=8),2):
                if rn in wanted:
                    chosen.append(tuple((c.value,c.data_type,c.number_format) for c in cells))
            if len(chosen)!=5000:
                raise ValueError('来源行缺失，停止重建。')
            selected[ws.title]=chosen
    finally:
        wb.close()
    output.parent.mkdir(parents=True,exist_ok=True)
    generated=output.with_name(f'{output.stem}.{uuid4().hex}.generated.xlsx')
    candidate=output.with_name(f'{output.stem}.{uuid4().hex}.verified.xlsx')
    try:
        result=Workbook()
        result.remove(result.active)
        result.epoch=epoch
        for name,rows in selected.items():
            ws=result.create_sheet(name)
            ws.append(HEADERS)
            for rn,cells in enumerate(rows,2):
                ws.append([c[0] for c in cells])
                for cn,(_,kind,fmt) in enumerate(cells,1):
                    if isinstance(fmt,str):
                        ws.cell(rn,cn).number_format=fmt
                    if kind=='f':
                        ws.cell(rn,cn).data_type='f'
        result.save(generated)
        result.close()
        restore_archive(generated,candidate,recipe)
        checked=0
        check=load_workbook(candidate,read_only=True,data_only=False)
        try:
            for ws in check.worksheets:
                if ws.max_row!=5001:
                    raise ValueError('重建行数不一致。')
                for expected,cells in zip(selected[ws.title],ws.iter_rows(min_row=2,max_col=8),strict=True):
                    for (value,kind,_),actual in zip(expected,cells,strict=True):
                        equal=(Decimal(str(value))==Decimal(str(actual.value)) if kind=='n' and value is not None
                               else value==actual.value)
                        if actual.data_type!=kind or not equal:
                            raise ValueError(f'读取值/类型不一致：{ws.title}/{actual.coordinate}')
                    checked+=1
        finally:
            check.close()
        if checked!=10000 or file_hash(source)!=manifest['source_sha256']:
            raise ValueError('重建期间源文件或行数变化。')
        if any(file_hash(evidence/name)!=h for name,h in preserved.items()):
            raise ValueError('重建期间清单变化。')
        if output.exists():
            raise ValueError('发布前已有样本出现，禁止覆盖。')
        candidate.rename(output)
        receipt={'rebuilt_at':datetime.now().astimezone().isoformat(),'scope':'Same fixed sample, no reselection',
                 'rows_verified':checked,'sample_sha256':file_hash(output),'source_sha256':manifest['source_sha256'],
                 'preserved_files':preserved,'script_sha256':file_hash(Path(__file__)),
                 'original_packaging_dates_restored':True,'work_time':'This receipt timestamp is the actual rebuild time'}
        target=evidence/f'rebuild-{uuid4().hex}.json'
        target.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
        print(json.dumps(receipt,ensure_ascii=False,indent=2))
        return receipt
    finally:
        for owned in [generated,candidate]:
            if owned.exists():
                owned.unlink()


if __name__=='__main__':
    rebuild_sample(ROOT/'data/raw/online_retail_II.xlsx',ROOT/'data/processed/performance-10000.xlsx',
                   ROOT/'evidence/performance/sample-v1')
