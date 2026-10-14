"""D-33 fixed whole-invoice sample; preserves source mappings and read values."""
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from decimal import Decimal
import csv
import json
from pathlib import Path
import sys
from uuid import uuid4
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'src'))
from vcr03.ingest.rules import Cell, HEADERS, key  # noqa: E402
from vcr03.ingest.service import file_hash  # noqa: E402

EXPECTED = 'bcbe73b35f5b7babf197fb0cb983a11f5d9ff929078d4aa53d171b1f2df2e980'
STARTS = {'Year 2009-2010': 131367, 'Year 2010-2011': 135479}


def build_sample():
    source = ROOT/'data/raw/online_retail_II.xlsx'
    output = ROOT/'data/processed/performance-10000.xlsx'
    evidence = ROOT/'evidence/performance/sample-v1'
    if output.exists():
        raise ValueError('固定样本已存在；先核验，不覆盖或重新选样。')
    if evidence.exists():
        from rebuild_performance_sample import rebuild_sample
        return rebuild_sample(source,output,evidence)
    assert file_hash(source) == EXPECTED
    wb = load_workbook(source, read_only=True, data_only=False)
    selected, reports = {}, []
    for ws in wb.worksheets:
        assert tuple(c.value for c in next(ws.iter_rows(max_row=1,max_col=8))) == HEADERS
        groups = defaultdict(list)
        for rn, cells in enumerate(ws.iter_rows(min_row=2,max_col=8), 2):
            inv, _, error = key(Cell(cells[0].value,cells[0].data_type))
            if error:
                raise ValueError('当前官方文件出现无法识别单号，先报告，不筛掉后继续。')
            groups[inv].append((rn,tuple((c.value,c.data_type,c.number_format) for c in cells)))
        remaining, chosen, invoice_count = 5000, [], 0
        for inv, lines in groups.items():
            if lines[0][0] >= STARTS[ws.title] and len(lines) <= remaining:
                chosen.extend(lines)
                remaining -= len(lines)
                invoice_count += 1
                if not remaining:
                    break
        if remaining:
            raise ValueError(f'{ws.title}整单不足5000行，剩余{remaining}；须确认，不截单。')
        chosen.sort(key=lambda x:x[0])
        selected[ws.title] = chosen
        reports.append({'sheet':ws.title,'threshold_excel_row':STARTS[ws.title],
                        'first_excel_row':chosen[0][0],'last_excel_row':chosen[-1][0],
                        'rows':len(chosen),'invoices':invoice_count,'whole_invoices_verified':True})
        print(json.dumps(reports[-1],ensure_ascii=False),flush=True)
    epoch = wb.epoch
    wb.close()
    result = Workbook()
    result.remove(result.active)
    result.epoch = epoch
    mapping = []
    for name, lines in selected.items():
        ws = result.create_sheet(name)
        ws.append(HEADERS)
        for sample_row,(original_row,cells) in enumerate(lines,2):
            ws.append([c[0] for c in cells])
            for column,(_,kind,fmt) in enumerate(cells,1):
                if isinstance(fmt,str):
                    ws.cell(sample_row,column).number_format = fmt
                if kind == 'f':
                    ws.cell(sample_row,column).data_type = 'f'
            mapping.append({'source_sha256':EXPECTED,'source_sheet':name,'source_excel_row':original_row,
                            'sample_sheet':name,'sample_excel_row':sample_row})
    output.parent.mkdir(parents=True,exist_ok=True)
    pending=output.with_name(f'{output.stem}.{uuid4().hex}.pending.xlsx')
    result.save(pending)
    result.close()
    check = load_workbook(pending,read_only=True,data_only=False)
    compared=0
    for ws in check.worksheets:
        assert ws.max_row == 5001
        for (_, expected), cells in zip(selected[ws.title], ws.iter_rows(min_row=2,max_col=8),strict=True):
            for (value,kind,_fmt), actual in zip(expected,cells,strict=True):
                assert actual.data_type == kind, (ws.title,actual.coordinate,kind,actual.data_type)
                if kind == 'n' and value is not None:
                    assert Decimal(str(value)) == Decimal(str(actual.value))
                else:
                    assert value == actual.value
            compared+=1
    check.close()
    assert compared == len(mapping) == 10000 and file_hash(source) == EXPECTED
    assert not output.exists()
    pending.rename(output)
    evidence.mkdir(parents=True)
    origin=evidence/'source_rows.csv'
    with origin.open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=list(mapping[0]))
        writer.writeheader()
        writer.writerows(mapping)
    manifest={'prepared_at':datetime.now(timezone(timedelta(hours=8))).isoformat(),
              'source_sha256':EXPECTED,'sample':'data/processed/performance-10000.xlsx',
              'sample_sha256':file_hash(output),'mapping':'evidence/performance/sample-v1/source_rows.csv',
              'mapping_sha256':file_hash(origin),'script_sha256':file_hash(Path(__file__)),
              'rules':'D-33; whole-invoice greedy; raw rows included before exclusions; fixed starts; no outcome filtering',
              'worksheets':reports,'rows':compared,'read_values_and_types_verified':True,
              'source_unchanged':True,'performance_acceptance':'not executed; A/B/initial-ticket pipeline required'}
    (evidence/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(manifest,ensure_ascii=False,indent=2))


if __name__ == '__main__':
    build_sample()
